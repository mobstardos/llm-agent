"""DOCX extractor."""
from __future__ import annotations

import logging
import time
from pathlib import Path

from src.extraction.base import ExtractionResult

logger = logging.getLogger(__name__)


class DocxExtractor:
    id = "docx"
    version = "0.1.0"
    mime_types = [
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ]
    priority = 100

    def can_handle(self, path: Path, mime: str) -> bool:
        return "wordprocessingml" in mime or path.suffix.lower() == ".docx"

    def extract(
        self, path: Path, max_chars: int = 200000, **kwargs,
    ) -> ExtractionResult:
        t0 = time.perf_counter()
        warnings: list[str] = []

        try:
            from docx import Document
        except ImportError:
            return ExtractionResult(
                source=str(path),
                mime="application/vnd.openxmlformats-officedocument."
                     "wordprocessingml.document",
                extractor_id=self.id, extractor_version=self.version,
                text="", warnings=["python-docx не установлен"],
            )

        try:
            doc = Document(str(path))
        except Exception as e:
            return ExtractionResult(
                source=str(path),
                mime="application/vnd.openxmlformats-officedocument."
                     "wordprocessingml.document",
                extractor_id=self.id, extractor_version=self.version,
                text="", warnings=[f"Ошибка чтения: {e}"],
            )

        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        text = "\n".join(paragraphs)

        # Таблицы
        tables: list[list[list]] = []
        for t in doc.tables:
            rows = []
            for row in t.rows:
                rows.append([c.text for c in row.cells])
            tables.append(rows)

        truncated = False
        if len(text) > max_chars:
            text = text[:max_chars]
            truncated = True
            warnings.append(f"Обрезано до {max_chars} символов")

        metadata = {
            "paragraphs": len(paragraphs),
            "tables": len(tables),
        }
        try:
            props = doc.core_properties
            metadata.update({
                "author": props.author or "",
                "title": props.title or "",
                "created": str(props.created) if props.created else "",
            })
        except Exception:
            pass

        duration_ms = (time.perf_counter() - t0) * 1000
        return ExtractionResult(
            source=str(path),
            mime="application/vnd.openxmlformats-officedocument."
                 "wordprocessingml.document",
            extractor_id=self.id, extractor_version=self.version,
            text=text, chars=len(text), tables=tables,
            metadata=metadata,
            truncated=truncated, warnings=warnings,
            duration_ms=duration_ms,
        )
