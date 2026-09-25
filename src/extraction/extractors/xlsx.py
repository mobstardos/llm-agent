"""XLSX extractor."""
from __future__ import annotations

import logging
import time
from pathlib import Path

from src.extraction.base import ExtractionResult

logger = logging.getLogger(__name__)


class XlsxExtractor:
    id = "xlsx"
    version = "0.1.0"
    mime_types = [
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ]
    priority = 100

    def can_handle(self, path: Path, mime: str) -> bool:
        return "spreadsheetml" in mime or path.suffix.lower() == ".xlsx"

    def extract(
        self, path: Path, max_rows_per_sheet: int = 100,
        max_chars: int = 200000, **kwargs,
    ) -> ExtractionResult:
        t0 = time.perf_counter()
        warnings: list[str] = []

        try:
            from openpyxl import load_workbook
        except ImportError:
            return ExtractionResult(
                source=str(path),
                mime="application/vnd.openxmlformats-officedocument."
                     "spreadsheetml.sheet",
                extractor_id=self.id, extractor_version=self.version,
                text="", warnings=["openpyxl не установлен"],
            )

        try:
            wb = load_workbook(str(path), data_only=False, read_only=True)
        except Exception as e:
            return ExtractionResult(
                source=str(path),
                mime="application/vnd.openxmlformats-officedocument."
                     "spreadsheetml.sheet",
                extractor_id=self.id, extractor_version=self.version,
                text="", warnings=[f"Ошибка чтения: {e}"],
            )

        parts: list[str] = []
        tables: list[list[list]] = []
        sheets_info: dict[str, dict] = {}
        truncated = False

        for name in wb.sheetnames:
            ws = wb[name]
            rows: list[list] = []
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= max_rows_per_sheet:
                    truncated = True
                    warnings.append(
                        f"Лист '{name}': обработаны первые {max_rows_per_sheet} строк"
                    )
                    break
                rows.append([c if c is not None else "" for c in row])

            tables.append(rows)
            sheets_info[name] = {
                "rows_shown": len(rows),
                "rows_total": ws.max_row or 0,
                "cols": ws.max_column or 0,
            }

            # Текстовое представление
            parts.append(f"## Лист: {name}")
            for r in rows[:50]:
                parts.append("\t".join(str(c) for c in r))

        text = "\n".join(parts)
        if len(text) > max_chars:
            text = text[:max_chars]
            truncated = True

        duration_ms = (time.perf_counter() - t0) * 1000
        return ExtractionResult(
            source=str(path),
            mime="application/vnd.openxmlformats-officedocument."
                 "spreadsheetml.sheet",
            extractor_id=self.id, extractor_version=self.version,
            text=text, chars=len(text), tables=tables,
            metadata={"sheets": sheets_info},
            truncated=truncated, warnings=warnings,
            duration_ms=duration_ms,
        )
